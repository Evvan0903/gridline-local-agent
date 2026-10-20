"""Auditable local exports. SQL is retained alongside values and explicit formulas."""

from datetime import datetime, date
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from xml.sax.saxutils import escape

import numpy as np
import pandas as pd

from .workspace import dump, fingerprint


def excel_value(value):
    if isinstance(value, (dict, list, tuple, set)):
        return str(value)
    if pd.isna(value):
        return None
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def export_excel(tables, path, formulas=None, *, financial=False):
    path = Path(path).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    sheets, names = [], set()
    for name, frame in tables.items():
        clean = re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "Table"
        if clean in names:
            raise ValueError("Excel sheet names collide after truncation")
        names.add(clean)
        rows = [list(frame.columns)] + [[excel_value(v) for v in row] for row in frame.itertuples(index=False, name=None)]
        if len(rows) > 1048576:
            raise ValueError("Excel row limit exceeded; export CSV instead")
        sheets.append({"name": clean, "rows": rows, "formulas": (formulas or {}).get(name, {})})
    assets = Path(__file__).parent / "assets"
    if (assets / "node_modules" / "@oai" / "artifact-tool").exists():
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "payload.json"
            source.write_text(dump({"sheets": sheets, "financial": financial}))
            subprocess.run(["node", str(assets / "export.mjs"), str(source), str(path)], check=True, timeout=300,
                           stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    else:
        # Portable fully-local fallback when the optional artifact runtime is not installed.
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        wb.remove(wb.active)
        for table in sheets:
            sheet = wb.create_sheet(table["name"])
            for row in table["rows"]:
                sheet.append(["'" + v if isinstance(v, str) and v.startswith("=") else v for v in row])
            for cell, value in table["formulas"].items():
                sheet[cell] = value
            for cell in sheet[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="243750")
            sheet.freeze_panes = "A2"
        wb.save(path)
    return {"path": str(path), "sha256": fingerprint(path), "sheets": [s["name"] for s in sheets]}


def export_pdf(title, sections, path):
    """Optional layer: bounded text/tables, page breaks set by the report author."""
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    styles = getSampleStyleSheet()
    path = Path(path).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 12)]
    for section in sections:
        if section.get("page_break"):
            story.append(PageBreak())
        story.append(Paragraph(escape(section["title"]), styles["Heading2"]))
        for text in section.get("paragraphs", []):
            story += [Paragraph(escape(str(text)), styles["BodyText"]), Spacer(1, 7)]
        if section.get("rows"):
            data = [[Paragraph(escape(str(v)), styles["BodyText"]) for v in row] for row in section["rows"]]
            widths = section.get("widths", [7.0 * inch / len(data[0])] * len(data[0]))
            table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#E7EDF4")),
                                       ("VALIGN", (0,0), (-1,-1), "TOP"),
                                       ("LINEBELOW", (0,0), (-1,0), .5, colors.grey),
                                       ("BOTTOMPADDING", (0,0), (-1,-1), 7)]))
            story.append(table)
    def footer(canvas, document):
        canvas.setFont("Helvetica", 8)
        canvas.drawString(36, 22, f"Local reproducible analysis | Page {document.page}")
    SimpleDocTemplate(str(path), pagesize=(612,792), leftMargin=36, rightMargin=36,
                      topMargin=30, bottomMargin=36).build(story, onFirstPage=footer, onLaterPages=footer)
    return {"path": str(path), "sha256": fingerprint(path)}


def export_query(workspace, query_id, name="analysis", excel=True, pdf=False, csv=True, kind="snapshot"):
    """Explicit result snapshot, NOT a dynamic financial workbook/model builder."""
    if not re.fullmatch(r"[a-zA-Z0-9_-]+", name):
        raise ValueError("Output name must contain only letters, digits, hyphens and underscores")
    if kind != "snapshot":
        raise ValueError("This tool exports result snapshots, not dynamic Excel models. A dynamic model requires explicit formulas, assumptions and refresh dependencies; no static substitute was created.")
    if not (csv or excel or pdf):
        raise ValueError("Select at least one export format")
    evidence = workspace.query_evidence(query_id)
    frame = workspace.query_result_frame(query_id)
    directory = workspace.root / "outputs" / name
    if directory.exists():
        raise ValueError("Export destination already exists; choose a new name to preserve existing files")
    directory.parent.mkdir(parents=True, exist_ok=True)
    temporary = Path(tempfile.mkdtemp(prefix=".export-", dir=directory.parent))
    outputs = {"status": "exported", "kind": "result_snapshot", "query_id": query_id,
               "dataset_versions": {t: m["version_id"] for t, m in evidence["dependencies"].items()},
               "lineage": str(directory / "lineage.json"), "report": str(directory / "report.md")}
    try:
        if csv:
            frame.to_csv(temporary / "result.csv", index=False)
            outputs["csv"] = str(directory / "result.csv")
        (temporary / "report.md").write_text(f"# {name}\n\nResult snapshot, not a dynamic model. Evidence: `{query_id}`. Full result rows: {len(frame)}.\n\n```sql\n{evidence['sql']}\n```\n\nParameters: `{dump(evidence['parameters'])}`\n\nSource versions, dependencies, hashes and transformations: `lineage.json`. No business meaning inferred. PDF tables show at most 15 rows.\n")
        if excel:
            output = export_excel({"Result": frame, "Provenance": pd.DataFrame(
                [{"evidence": query_id, "sql": evidence["sql"], "parameters": dump(evidence["parameters"]),
                  "dependencies": dump(evidence["dependencies"]), "result_digest": dump(evidence["result_digest"])}])}, temporary / "analysis.xlsx")
            outputs["excel"] = {**output, "path": str(directory / "analysis.xlsx")}
        if pdf:
            output = export_pdf(name, [{"title": "Result snapshot — calculation evidence", "paragraphs": [query_id, evidence["sql"], dump(evidence["parameters"]),
                f"Full result rows: {len(frame)}. Preview below: at most 15 rows. Complete lineage in lineage.json."],
                "rows": [list(frame.columns)] + frame.head(15).astype(str).values.tolist()}], temporary / "report.pdf")
            outputs["pdf"] = {**output, "path": str(directory / "report.pdf")}
        workspace.query_evidence(query_id)
        evidence["export"] = {"kind": "result_snapshot", "row_count": len(frame),
            "files_sha256": {p.name: fingerprint(p) for p in temporary.iterdir()}}
        (temporary / "lineage.json").write_text(dump(evidence))
        temporary.rename(directory)
    finally:
        if temporary.exists():
            shutil.rmtree(temporary)
    workspace.record("export", query_id=query_id, outputs=outputs, result_digest=evidence["result_digest"])
    return outputs
