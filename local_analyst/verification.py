"""Local artifact verification. Native recalculation is independent of model prose."""
from pathlib import Path
import json,math,re,subprocess,hashlib,uuid
from openpyxl import load_workbook

def recalculate_workbook(path,destination,changes):
    path=Path(path).resolve();destination=Path(destination).resolve()
    if destination.exists():raise ValueError('Validation destination exists')
    before=hashlib.sha256(path.read_bytes()).hexdigest();book=load_workbook(path,data_only=False)
    if 'Parameters' not in book:raise ValueError('No parameter sheet; cannot prove dynamic dependencies')
    allowed={f'B{r}' for r in range(2,book['Parameters'].max_row+1)}
    if set(changes)-allowed or any(type(v) not in (int,float) or not math.isfinite(v) for v in changes.values()):raise ValueError('Only finite numeric changes to declared parameter cells allowed')
    destination.parent.mkdir(parents=True,exist_ok=True)
    script=Path(__file__).parent/'assets/verify_model.mjs'
    subprocess.run(['node',str(script),str(path),str(destination),json.dumps(changes)],check=True,timeout=90,stdout=subprocess.PIPE,stderr=subprocess.PIPE)
    after=load_workbook(destination,data_only=True)
    errors=[f'{s.title}!{c.coordinate}' for s in after for row in s for c in row if c.data_type=='e']
    if hashlib.sha256(path.read_bytes()).hexdigest()!=before:raise ValueError('Original artifact changed')
    return {'status':'PASS' if not errors else 'FAIL','errors':errors,'original_sha256':before,'verification_copy':str(destination),'changes':changes,
         'summary':list(after['Summary'].values) if 'Summary' in after else [],'limitations':['Recalculation alone is not an independent business oracle']}
