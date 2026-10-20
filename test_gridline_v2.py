"""Offline acceptance: SDK MockTransport, real Agent/tools, restricted compiler and guards."""
import copy
import json
from pathlib import Path
from decimal import Decimal
import tempfile
import unittest
from unittest.mock import Mock, patch
import httpx
import pandas as pd
from openai import OpenAI
from fastapi.testclient import TestClient
from local_analyst.providers import SpendLedger, OpenAIProvider, ProviderBlocked, estimated_cost, response_tools
from local_analyst.agent import AnalystAgent
from local_analyst.financial_model import evaluate, build_financial_model, expression
from local_analyst.workspace import Workspace
from web_api.app import create_app


def specification():
    return {'version':1,'rule_ref':'synthetic unit price reviewed','source_key':'id',
      'parameters':{'multiplier':{'value':'1.2'}},'calculations':[
      {'id':'base','expression':{'mul':[{'field':'amount'},'0.1']}},
      {'id':'scenario','expression':{'mul':[{'field':'base'},{'param':'multiplier'}]}},
      {'id':'delta','expression':{'sub':[{'field':'scenario'},{'field':'base'}]}}],'totals':['base','scenario','delta']}


def response(output, status='completed', usage=None):
    return {'id':'resp_mock','object':'response','created_at':1,'model':'gpt-6-luna','status':status,
        'output':output,'usage':usage or {'input_tokens':100,'output_tokens':20,'total_tokens':120,
        'input_tokens_details':{'cached_tokens':10},'output_tokens_details':{'reasoning_tokens':5}}}


class ProviderTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.ledger=SpendLedger(self.root/'spend.sqlite3');self.addCleanup(self.ledger.db.close)
        self.approval=self.ledger.approve('session','scope','gpt-6-luna',data_authorized=True,paid_authorized=True)
        self.requests=[];self.responses=[]
        def send(request):
            self.requests.append(json.loads(request.content));r=self.responses.pop(0)
            return httpx.Response(r[0],json=r[1]) if isinstance(r,tuple) else httpx.Response(200,json=r)
        self.client=OpenAI(api_key='dummy-offline-only',http_client=httpx.Client(transport=httpx.MockTransport(send)),max_retries=0)
        self.addCleanup(self.client.close)
        self.provider=OpenAIProvider(self.ledger,session='session',scope='scope',approval=self.approval,run='run',client=self.client)
    def test_official_sdk_reasoning_continuity_call_ids_and_multiturn(self):
        reasoning={'id':'rs_1','type':'reasoning','summary':[],'encrypted_content':'opaque_mock'}
        call={'type':'function_call','id':'fc_1','call_id':'real_call_1','name':'inspect_schema','arguments':'{}','status':'completed'}
        self.responses=[response([reasoning,call]),response([{'type':'message','id':'msg_1','role':'assistant','status':'completed','content':[{'type':'output_text','text':'done','annotations':[]}]}])]
        history=[{'role':'system','content':'Use tools.'},{'role':'user','content':'Inspect'}]
        first,_=self.provider.complete('gpt-6-luna',history,[])
        history.extend([first,{'role':'tool','tool_call_id':'real_call_1','content':'{"ok":true}'},{'role':'user','content':'Continue'}])
        second,meta=self.provider.complete('gpt-6-luna',history,[])
        self.assertEqual(second['content'],'done');self.assertEqual(first['tool_calls'][0]['id'],'real_call_1')
        self.assertIn(reasoning,self.requests[1]['input']);self.assertIn(call,self.requests[1]['input'])
        self.assertIn({'type':'function_call_output','call_id':'real_call_1','output':'{"ok":true}'},self.requests[1]['input'])
        self.assertFalse(self.requests[0]['store']);self.assertEqual(self.requests[0]['service_tier'],'default')
        self.assertFalse(self.requests[0]['parallel_tool_calls']);self.assertEqual(meta['spend']['models']['gpt-6-luna']['requests'],2)
    def test_confidential_on_blocks_before_sdk_or_reservation(self):
        self.provider.confidential_mode=True
        with self.assertRaisesRegex(ProviderBlocked,"CONFIDENTIAL_MODE_UNAVAILABLE"):
            self.provider.complete("gpt-6-luna",[],[])
        self.assertEqual(self.requests,[])
        self.assertEqual(self.ledger.summary()["models"]["gpt-6-luna"]["requests"],0)

    def test_application_secrets_removed_but_authorized_records_preserved(self):
        self.responses=[response([])]
        with patch.dict("os.environ",{"GRIDLINE_APP_SECRET":"offline-secret-sentinel"}):
            self.provider.complete("gpt-6-luna",[{"role":"user","content":"Employee Example; amount 100; offline-secret-sentinel"}],[])
        text=json.dumps(self.requests)
        self.assertNotIn("offline-secret-sentinel",text)
        self.assertIn("Employee Example",text)

    def test_multiple_calls_preserved(self):
        self.responses=[response([{'type':'function_call','call_id':f'c{i}','name':'inspect_schema','arguments':'{}'} for i in range(2)])]
        m,_=self.provider.complete('gpt-6-luna',[{'role':'user','content':'inspect'}],[])
        self.assertEqual([c['id'] for c in m['tool_calls']],['c0','c1'])
    def test_flattened_non_strict_open_map_schema(self):
        w=Workspace(self.root/'w');self.addCleanup(w.db.close);a=AnalystAgent(w)
        tools=response_tools(a.tools)
        self.assertTrue(all(t['type']=='function' and 'function' not in t and t['strict'] is False for t in tools))
        self.assertEqual(next(t for t in tools if t['name']=='query_sql')['parameters']['required'],['sql'])
    def test_no_approval_never_requests(self):
        self.provider.approval='absent'
        with self.assertRaisesRegex(ProviderBlocked,'AUTHORIZATION'):self.provider.complete('gpt-6-luna',[],[])
        self.assertEqual(self.requests,[])
    def test_model_escalation_never_requests(self):
        with self.assertRaisesRegex(ProviderBlocked,'AUTHORIZATION'):self.provider.complete('gpt-6.1-sol',[],[])
        self.assertEqual(self.requests,[])
    def test_source_scope_mismatch_never_requests(self):
        self.provider.scope='new_source'
        with self.assertRaises(ProviderBlocked):self.provider.complete('gpt-6-luna',[],[])
        self.assertEqual(self.requests,[])
    def test_invalid_credential_one_attempt_no_exception_body(self):
        self.responses=[(401,{'error':{'message':'sensitive key echoed here','type':'invalid_request_error','code':'invalid_api_key'}})]
        with self.assertRaisesRegex(ProviderBlocked,'AuthenticationError') as caught:self.provider.complete('gpt-6-luna',[],[])
        self.assertNotIn('sensitive',str(caught.exception));self.assertEqual(len(self.requests),1)
        s=self.ledger.summary()['models']['gpt-6-luna'];self.assertEqual(s['unknown_usage_requests'],1)
        with self.assertRaisesRegex(ProviderBlocked,'UNRECONCILED'):self.provider.complete('gpt-6-luna',[],[])
        self.assertEqual(len(self.requests),1)
    def test_incomplete_response_charged_never_success(self):
        self.responses=[response([],status='incomplete')]
        with self.assertRaisesRegex(ProviderBlocked,'INCOMPLETE'):self.provider.complete('gpt-6-luna',[],[])
        self.assertEqual(self.ledger.summary()['models']['gpt-6-luna']['output_tokens'],20)
    def test_budget_rejection_is_before_request(self):
        self.provider.approval=self.ledger.approve('session','scope','gpt-6-luna',data_authorized=True,paid_authorized=True,run_cap='0.000001')
        with self.assertRaisesRegex(ProviderBlocked,'BUDGET_EXCEEDED'):self.provider.complete('gpt-6-luna',[],[])
        self.assertEqual(self.requests,[])
    def test_request_limit_and_persistence(self):
        aid=self.ledger.approve('session','scope','gpt-6-luna',data_authorized=True,paid_authorized=True,max_requests=1)
        k=self.ledger.reserve(aid,'session','scope','bounded','gpt-6-luna',Decimal('.001'))
        self.ledger.finish(k,{'input_tokens':100,'output_tokens':20})
        with self.assertRaisesRegex(ProviderBlocked,'REQUEST_LIMIT'):self.ledger.reserve(aid,'session','scope','bounded','gpt-6-luna',Decimal('.001'))
        other=SpendLedger(self.root/'spend.sqlite3');self.addCleanup(other.db.close)
        self.assertEqual(other.summary()['models']['gpt-6-luna']['requests'],1)
    def test_cancelled_no_request(self):
        self.provider.cancelled=lambda:True
        with self.assertRaisesRegex(ProviderBlocked,'CANCELLED'):self.provider.complete('gpt-6-luna',[],[])
        self.assertEqual(self.requests,[])
    def test_cached_reasoning_accounted_without_double_count(self):
        self.assertEqual(estimated_cost('gpt-6.1-sol',1000,100,cached=200),Decimal('0.00262'))
        with self.assertRaises(ProviderBlocked):estimated_cost('gpt-6-luna',10,5,cached=11)
    def test_key_redacted_from_sent_content(self):
        self.responses=[response([])]
        with patch.dict('os.environ',{'OPENAI_API_KEY':'dummy-secret-value'}):
            self.provider.complete('gpt-6-luna',[{'role':'user','content':'dummy-secret-value'}],[])
        self.assertNotIn('dummy-secret-value',json.dumps(self.requests));self.assertIn('REDACTED',json.dumps(self.requests))

    def test_real_agent_sql_over_official_sdk_mock_transport(self):
        path=self.root/'source.csv';path.write_text('id,amount\n1,10.10\n2,20.20\n')
        w=Workspace(self.root/'agent');self.addCleanup(w.db.close);table=w.load(path)['tables'][0]['table_id']
        self.responses=[response([{'type':'function_call','id':'fc_sql','call_id':'sdk_sql_1','name':'query_sql','arguments':json.dumps({'sql':f'SELECT SUM(amount) AS total FROM {table}'})}]),response([{'type':'message','id':'msg_done','role':'assistant','status':'completed','content':[{'type':'output_text','text':'The total is 30.30.','annotations':[]}]}])]
        a=AnalystAgent(w,model='gpt-6-luna',provider=self.provider)
        answer=a.ask('Execute SQL to calculate the total amount.')
        self.assertEqual(answer['status'],'answered',answer)
        self.assertEqual(answer['trace'][0]['result']['rows'],[[Decimal('30.30')]])
        output=next(x for x in self.requests[1]['input'] if x.get('type')=='function_call_output')
        self.assertEqual(output['call_id'],'sdk_sql_1');self.assertIn('30.30',output['output'])
        self.assertEqual(self.ledger.summary()['models']['gpt-6-luna']['requests'],2)

    def test_timeout_no_retry_retains_reservation(self):
        def fail(request):
            self.requests.append({})
            raise httpx.ReadTimeout('dummy timeout')
        client=OpenAI(api_key='dummy',http_client=httpx.Client(transport=httpx.MockTransport(fail)),max_retries=0)
        self.addCleanup(client.close);self.provider.client=client
        with self.assertRaisesRegex(ProviderBlocked,'APITimeoutError'):self.provider.complete('gpt-6-luna',[],[])
        self.assertEqual(len(self.requests),1)
        self.assertGreater(Decimal(self.ledger.summary()['models']['gpt-6-luna']['committed_usd']),0)

    def test_separate_sol_cap_and_no_run_reset_bypass(self):
        a=self.ledger.approve('s','case','gpt-6.1-sol',data_authorized=True,paid_authorized=True,run_cap='2',sol_cap='1')
        k=self.ledger.reserve(a,'s','case','one','gpt-6.1-sol',Decimal('.7'))
        self.ledger.finish(k,{'input_tokens':100000,'output_tokens':50000})
        with self.assertRaisesRegex(ProviderBlocked,'BUDGET_EXCEEDED'):
            self.ledger.reserve(a,'s','case','two','gpt-6.1-sol',Decimal('.4'))

    def test_cache_write_usage_uses_sdk_field_and_not_cached_price(self):
        k=self.ledger.reserve(self.approval,'session','scope','write','gpt-6-luna',Decimal('.01'))
        self.ledger.finish(k,{'input_tokens':1000,'output_tokens':100,'input_tokens_details':{'cached_tokens':200,'cache_write_tokens':300}})
        x=self.ledger.summary()['models']['gpt-6-luna']
        self.assertEqual(Decimal(x['observed_estimate_usd']),Decimal('.0001395'))
        self.assertEqual(x['cache_write_tokens'],300)

    def test_orphaned_reservation_blocks_after_restart(self):
        self.ledger.reserve(self.approval,'session','scope','old','gpt-6-luna',Decimal('.01'))
        other=SpendLedger(self.root/'spend.sqlite3');self.addCleanup(other.db.close)
        with self.assertRaisesRegex(ProviderBlocked,'UNRECONCILED'):
            other.reserve(self.approval,'session','scope','new','gpt-6-luna',Decimal('.01'))


class FormulaTests(unittest.TestCase):
    def setUp(self):
        self.spec=specification();self.frame=pd.DataFrame({'id':['a','b'],'amount':[100,200]})
    def test_synthetic_changed_parameter_propagates_independent_oracle(self):
        x=evaluate(self.spec,self.frame,{'multiplier':'1.5'})
        self.assertEqual(x.scenario.tolist(),[Decimal('15'),Decimal('30')]);self.assertEqual(sum(x.delta),Decimal('15'))
    def test_arbitrary_formula_and_forward_refs_denied(self):
        for expr in ({'formula':'=WEBSERVICE("x")'},{'field':'missing'},{'div':[1,0]}):
            s=copy.deepcopy(self.spec);s['calculations'][0]['expression']=expr
            with self.assertRaises((ValueError,ArithmeticError)):evaluate(s,self.frame)
    def test_missing_period_and_unapproved_assumptions_block(self):
        for extra in ({'approved_result':True,'coverage_complete':False},{'assumptions':[{'status':'requires_approval'}]}):
            s={**self.spec,**extra}
            with self.assertRaises(ValueError):evaluate(s,self.frame)
    def test_duplicate_or_missing_source_keys_rejected(self):
        for ids in (['a','a'],['a',None]):
            with self.assertRaises(ValueError):evaluate(self.spec,pd.DataFrame({'id':ids,'amount':[1,2]}))
    def test_rounding_negative_half_up_and_zero_division(self):
        self.assertEqual(expression({'round':['-1.005',2]}, {},{},'decimal'),Decimal('-1.01'))
    def test_reopen_formula_values_and_destination_preservation(self):
        with tempfile.TemporaryDirectory() as d:
            dest=Path(d)/'model'
            result=build_financial_model(self.spec,self.frame,dest,{'source':'synthetic'})
            self.assertEqual(result['validation']['totals']['scenario'],'36.00')
            from openpyxl import load_workbook
            self.assertEqual(load_workbook(dest/'model.xlsx',data_only=True)['Summary']['B4'].value,6)
            self.assertTrue(load_workbook(dest/'model.xlsx')['Summary']['B3'].value.startswith('=SUM'))
            with self.assertRaises(ValueError):build_financial_model(self.spec,self.frame,dest,{})

    def test_scripted_agent_current_sql_to_reviewed_dynamic_artifact(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);source=root/'source.csv';source.write_text('id,amount\na,100\nb,200\n')
            w=Workspace(root/'workspace');self.addCleanup(w.db.close);table=w.load(source)['tables'][0]['table_id']
            plan=root/'plan.json';plan.write_text(json.dumps({'plan_id':'synthetic_financial','financial_model':self.spec}))
            plan_id=w.register_plan(plan);agent=AnalystAgent(w,model='scripted-local')
            calls=[lambda:{'role':'assistant','content':'','tool_calls':[{'function':{'name':'query_sql','arguments':{'sql':f'SELECT id,amount FROM {table}'}}}]},
              lambda:{'role':'assistant','content':'','tool_calls':[{'function':{'name':'build_financial_model','arguments':{'plan_id':plan_id,'query_id':agent._query_ids[-1],'name':'reviewed_model'}}}]},
              lambda:{'role':'assistant','content':'Created the requested dynamic workbook. Baseline 30, scenario 36 and delta 6.'}]
            def model(*args,**kwargs):
                out=Mock();out.json.return_value={'message':calls.pop(0)()};return out
            with patch('local_analyst.agent.requests.post',side_effect=model):
                result=agent.ask('Execute SQL over the registered inputs. Create a dynamic Excel model using the registered financial specification.')
            self.assertEqual(result['status'],'answered',result)
            artifact=result['trace'][1]['result'];self.assertEqual(artifact['kind'],'dynamic_financial_model')
            self.assertEqual(artifact['validation']['totals']['scenario'],'36.00')
            self.assertTrue(Path(artifact['path']).is_file())


class CloudHTTPTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name);self.app=create_app(self.root/'state',[self.root]);self.m=self.app.state.manager
        self.m.models=lambda:{'models':[]}
        self.client=TestClient(self.app,base_url='http://localhost');self.client.__enter__();self.addCleanup(lambda:self.client.__exit__(None,None,None))
        b=self.client.get('/api/bootstrap').json();self.client.headers['x-csrf-token']=b['csrf_token']
        self.w=self.client.post('/api/workspaces',json={'name':'w'}).json()['id']
        self.c=self.client.post(f'/api/workspaces/{self.w}/conversations',json={}).json()['id']
    def approve(self,**extra):
        return self.client.post(f'/api/conversations/{self.c}/cloud-approval',json={'model':'gpt-6-luna','paid_authorized':True,'data_authorized':True,**extra})
    def test_both_authorizations_required_and_sol_gated(self):
        self.assertEqual(self.approve(data_authorized=False).status_code,403)
        self.assertEqual(self.approve(paid_authorized=False).status_code,403)
        self.assertEqual(self.approve(model='gpt-6.1-sol').status_code,409)
        self.assertEqual(self.m.spend.summary()['models']['gpt-6.1-sol']['requests'],0)
    def test_new_source_invalidates_consent(self):
        a=self.approve().json()
        self.client.post(f'/api/workspaces/{self.w}/upload',files={'file':('x.csv',b'id,amount\n1,10\n')})
        with patch.dict('os.environ',{'OPENAI_API_KEY':'dummy'}),patch('openai.OpenAI') as sdk:
            r=self.client.post(f'/api/conversations/{self.c}/messages',json={'content':'calculate','provider':'openai','model':'gpt-6-luna','approval_id':a['id']})
        self.assertEqual(r.status_code,403);sdk.assert_not_called()
    def test_cloud_attempt_missing_approval_no_network(self):
        with patch('openai.OpenAI') as sdk:
            r=self.client.post(f'/api/conversations/{self.c}/messages',json={'content':'hello','provider':'openai','model':'gpt-6-luna'})
        self.assertEqual(r.status_code,403);sdk.assert_not_called()
    def test_confidential_default_off_and_server_on_cannot_be_overridden(self):
        self.assertFalse(__import__("web_api.app",fromlist=["SettingsInput"]).SettingsInput().confidential_mode)
        self.assertEqual(self.approve(confidential_mode=True).status_code,409)
        self.assertEqual(self.approve().status_code,200)
        self.client.put("/api/settings",json={"confidential_mode":True})
        self.assertEqual(self.approve(confidential_mode=False).status_code,409)
        self.assertEqual(self.m.spend.summary()["models"]["gpt-6-luna"]["requests"],0)

    def test_spending_zero_is_observed_and_invoice_unavailable(self):
        x=self.client.get('/api/spending').json()
        self.assertEqual(x['models']['gpt-6.1-sol']['requests'],0);self.assertIsNone(x['invoice_cost_usd'])

    def test_benchmark_cloud_not_run_no_accidental_sdk_requests(self):
        import time
        fixture=Path(__file__).parent/'reports/validation/web_ui/fixtures/Demo Sales.xlsx'
        f=self.client.post(f'/api/workspaces/{self.w}/upload',files={'file':('synthetic.xlsx',fixture.read_bytes())})
        self.assertEqual(f.status_code,200)
        source=self.client.get(f'/api/workspaces/{self.w}').json()['files'][0]['id']
        with patch('openai.OpenAI') as sdk:
            b=self.client.post('/api/benchmarks',json={'workspace_id':self.w,'source_id':source,'models':['gpt-6-luna','gpt-6.1-sol']})
            self.assertEqual(b.status_code,202,b.text)
            for _ in range(100):
                r=self.client.get('/api/benchmarks/'+b.json()['id']).json()
                if r['status']!='running':break
                time.sleep(.01)
            self.assertEqual([x['status'] for x in r['results']],['NOT RUN','NOT RUN'])
            sdk.assert_not_called()
        download=self.client.get('/api/benchmarks/'+b.json()['id']+'/report')
        self.assertEqual(download.status_code,200);self.assertIn('attachment',download.headers['content-disposition'])

    def test_benchmark_source_ownership_and_partial_consent_denied(self):
        fixture=Path(__file__).parent/'reports/validation/web_ui/fixtures/Demo Sales.xlsx'
        self.client.post(f'/api/workspaces/{self.w}/upload',files={'file':('synthetic.xlsx',fixture.read_bytes())})
        source=self.client.get(f'/api/workspaces/{self.w}').json()['files'][0]['id']
        other=self.client.post('/api/workspaces',json={'name':'other'}).json()['id']
        body={'workspace_id':other,'source_id':source,'models':['gpt-6-luna']}
        self.assertEqual(self.client.post('/api/benchmarks',json=body).status_code,403)
        body.update(workspace_id=self.w,paid_authorized=True)
        self.assertEqual(self.client.post('/api/benchmarks',json=body).status_code,403)

if __name__=='__main__':unittest.main()
