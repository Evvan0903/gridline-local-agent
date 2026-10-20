"""Frozen synthetic generalization task and independent source/Decimal acceptance.

No benchmark-specific solver is exposed to the model. Only user rule text and
sources are visible; expected values stay in the verifier after execution.
"""
from pathlib import Path
from decimal import Decimal as D,ROUND_HALF_UP
import json
import pandas as pd
from local_analyst.exporting import export_excel
from local_analyst.v3 import execute
from local_analyst.workspace import fingerprint

RULES='''June 2026 earned orders only. Canonicalize identical duplicate customer mappings before joining; do not multiply revenue. Missing FX requires a documented hold, not an invented rate or zero. Commission is 10% of eligible net USD where FX is available; local units = USD commission times units_per_usd. Scenario changes commission to 12.5% (editable rate parameter). Round each output to cents. Never sum different local currencies. Source data is synthetic. These written rules authorize a conditional model proposal, not payroll approval.'''
OBJECTIVES=[{'id':'plan','requirement':'Create an evidence-backed structured plan','kind':'plan'},
 {'id':'calculation','requirement':'Calculate June commission by eligible order in USD/local, identify join grain and dates','kind':'financial'},
 {'id':'exceptions','requirement':'Identify duplicate customer mapping and missing FX hold with source evidence','kind':'financial'},
 {'id':'excel','requirement':'Generate a dynamic Excel scenario with real parameter-linked formulas and reconciled totals','kind':'artifact','required_artifacts':['excel']},
 {'id':'report','requirement':'Generate a PDF conclusion including assumptions and missing FX, cite evidence','kind':'artifact','required_artifacts':['pdf']}]
TASK='''Analyze the attached synthetic financial sources using the General Analyst profile. Follow governing rules. Address original objective IDs plan/calculation/exceptions/excel/report. Inspect actual schemas; execute read-only SQL and check_join on the customer key. Propose a structured task plan. Calculate eligible June commission, preserve order_id as unique source key and exclude held/missing-FX rows from payable outputs. Document duplicate mappings and FX holds. Create a dynamic Excel financial model and a PDF financial conclusion. For independent machine verification, the calculation query and model Build must expose order_id, commission_usd, commission_local, scenario_usd, scenario_local, delta_usd. Use the generic define_financial_plan then build_financial_model tools; return their real IDs in separate dependent turns. Parameters should expose scenario_rate. Do not invent data or payroll approval. Reuse schema metadata; batch independent calls, preserve final synthesis within eight model requests.'''

def create(root):
    root=Path(root);root.mkdir(parents=True,exist_ok=False)
    orders=pd.DataFrame([['o1','cA','2026-06-05',120,'earned'],['o2','cB','2026-06-12',240,'earned'],['o3','cA','2026-05-31',90,'earned'],['o4','cA','2026-06-19',50,'refunded'],['o5','cC','2026-06-20',80,'earned']],columns=['order_id','customer_id','order_date','net_usd','status'])
    export_excel({'Orders':orders},root/'Orders.xlsx')
    pd.DataFrame([['cA','Synthetic A','USD'],['cA','Synthetic A','USD'],['cB','Synthetic B','EUR'],['cC','Synthetic C','JPY']],columns=['customer_id','customer_name','currency']).to_csv(root/'Customers.csv',index=False)
    pd.DataFrame([['USD',1],['EUR',.9]],columns=['currency','units_per_usd']).to_csv(root/'FX.csv',index=False)
    (root/'rules.txt').write_text(RULES);(root/'task.txt').write_text(TASK)
    return [root/'Orders.xlsx',root/'Customers.csv',root/'FX.csv']

def expected(sources,scenario_rate='.125'):
    orders=pd.read_excel(sources[0]);customers=pd.read_csv(sources[1]).drop_duplicates();fx=pd.read_csv(sources[2])
    assert customers.customer_id.is_unique
    orders['order_date']=pd.to_datetime(orders.order_date)
    relevant=orders[orders.status.eq('earned')&orders.order_date.ge('2026-06-01')&orders.order_date.lt('2026-07-01')]
    merged=relevant.merge(customers,on='customer_id',validate='many_to_one').merge(fx,on='currency',how='left',validate='many_to_one')
    rows={}
    for row in merged[merged.units_per_usd.notna()].itertuples():
        amount=D(str(row.net_usd));rate=D(str(row.units_per_usd));q=lambda x:x.quantize(D('.01'),rounding=ROUND_HALF_UP)
        b=amount*D('.10');s=amount*D(scenario_rate)
        rows[row.order_id]={'commission_usd':q(b),'commission_local':q(b*rate),'scenario_usd':q(s),'scenario_local':q(s*rate),'delta_usd':q(s)-q(b)}
    return rows

def verifier(sources):
    oracle=expected(sources)
    def check(w,answer,events):
        results=[]
        def add(id,status,reason):results.append({'objective_id':id,'status':status,'reason':reason})
        add('plan','PASS' if any(e.get('tool')=='set_task_plan' for e in events) or any(e.get('role')=='planner' and e['kind']=='model_response' for e in events) else 'FAIL','Observed model planning handoff; detailed plan quality requires separate review')
        matched=False;bad=False
        for evidence in w.queries.values():
            frame=w.query_result_frame(next(k for k,v in w.queries.items() if v is evidence))
            if not {'order_id',*next(iter(oracle.values()))}.issubset(frame.columns):continue
            if len(frame)!=len(oracle) or frame.order_id.duplicated().any():bad=True;continue
            try:
                correct=all(all(abs(D(str(row[c]))-values[c])<D('.005') for c in values) for k,values in oracle.items() for _,row in frame[frame.order_id.eq(k)].iterrows()) and set(frame.order_id)==set(oracle)
            except Exception:correct=False
            matched|=correct;bad|=not correct
        joins=[e for e in events if e.get('tool')=='check_join' and e.get('result',{}).get('status')!='error']
        add('calculation','PASS' if matched and joins else 'FAIL','Independent per-order Decimal/date/join-grain oracle matched' if matched and joins else 'Missing or incorrect source-backed per-order commission/currency query and explicit join evidence')
        text=(answer.get('answer','')+' '+json.dumps([e.get('result') for e in events if e.get('tool') in {'check_quality','query_sql'}],default=str)).lower()
        # Require actual investigative evidence, not simply echoing the written task.
        customer_table=next(t for t,m in w.catalog.items() if Path(m.get('source_path','')).name==Path(sources[1]).name)
        duplicate=any(e.get('tool')=='check_quality' and e.get('arguments',{}).get('table_id')==customer_table and e.get('arguments',{}).get('kind')=='duplicates' and e.get('arguments',{}).get('columns')==['customer_id'] and e.get('result',{}).get('affected_rows')==2 for e in events)
        # A row ID whose numeric value is 2 is not a duplicate count. For SQL,
        # require a real single-source COUNT over customer_id, or actual quality evidence.
        import sqlglot
        from sqlglot import exp
        for event in events:
            qid=event.get('result',{}).get('query_id');q=w.queries.get(qid)
            if event.get('tool')!='query_sql' or not q or set(q['dependencies'])!={customer_table}:continue
            tree=sqlglot.parse_one(q['sql'],read='duckdb')
            if not tree.find(exp.Count) or not tree.args.get('group'):continue
            data=event['result'];cols=data.get('columns',[])
            if 'customer_id' not in cols:continue
            for row in data.get('rows',[]):
                if row[cols.index('customer_id')]=='cA' and any(isinstance(v,(int,float)) and v==2 for i,v in enumerate(row) if i!=cols.index('customer_id')):duplicate=True
        missing=any(e.get('tool')=='query_sql' and any('o5' in row and 'JPY' in row for row in e.get('result',{}).get('rows',[])) for e in events)
        add('exceptions','PASS' if duplicate and missing and any(x in text for x in ('hold','missing','jpy')) else 'FAIL','Structural duplicate and missing-FX investigation evidence; no financial loss inferred from row counts')
        from openpyxl import load_workbook
        files=list((w.root/'outputs').rglob('*.xlsx'));valid=False
        for file in files:
            b=load_workbook(file,data_only=True);f=load_workbook(file,data_only=False)
            if 'Build' not in b or 'Parameters' not in b:continue
            headers=[c.value for c in b['Build'][1]]
            if not {'order_id',*next(iter(oracle.values()))}.issubset(headers):continue
            rows={str(row[0]):dict(zip(headers,row)) for row in b['Build'].iter_rows(min_row=2,values_only=True)}
            try:valid=set(rows)==set(oracle) and all(abs(D(str(rows[k][c]))-v)<D('.005') for k,vals in oracle.items() for c,v in vals.items()) and any(c.data_type=='f' for row in f['Build'] for c in row) and any('Parameters' in str(c.value) for row in f['Build'] for c in row)
            except Exception:valid=False
        if valid:
            from local_analyst.verification import recalculate_workbook
            import uuid
            params=load_workbook(file,data_only=True)['Parameters']
            cell=next((f'B{r}' for r in range(2,params.max_row+1) if params[f'A{r}'].value=='scenario_rate'),None)
            if cell:
                validation_copy=w.root/'validations'/('independent_'+uuid.uuid4().hex+'.xlsx')
                recalculate_workbook(file,validation_copy,{cell:.20})
                changed=load_workbook(validation_copy,data_only=True)['Build'];head=[c.value for c in changed[1]];lookup={str(row[0]):dict(zip(head,row)) for row in changed.iter_rows(min_row=2,values_only=True)}
                changed_oracle=expected(sources,'.20')
                valid=all(abs(D(str(lookup[k][c]))-v)<D('.005') for k,vals in changed_oracle.items() for c,v in vals.items())
            else:valid=False
        add('excel','PASS' if valid else 'FAIL','Saved formula workbook/detail oracle matched; native changed scenario-rate .20 recalculation verified independently in a copy' if valid else 'No qualifying dynamic workbook with independently correct per-order values')
        from pypdf import PdfReader
        pdfs=list((w.root/'outputs').rglob('*.pdf'));pdf_ok=False
        for file in pdfs:
            doc=PdfReader(file);t=' '.join(p.extract_text() or '' for p in doc.pages).lower()
            pdf_ok|=len(doc.pages)<=3 and '36' in t and '45' in t and any(x in t for x in ('missing','hold','jpy'))
        add('report','PASS' if pdf_ok and matched else 'FAIL','PDF/page/financial-figure/missing-FX checks; arbitrary prose correctness not universally proved')
        return results
    return check

def run_suite(root):
    root=Path(root);root.mkdir(parents=True,exist_ok=False);sources=create(root/'sources')
    common=dict(sources=sources,objectives=OBJECTIVES,rule_text=RULES,formats=['excel','pdf'],verifier=verifier(sources),max_requests=8,max_tools=24,max_seconds=600)
    results={}
    frozen={str(p.resolve()):fingerprint(p) for p in sources}
    (root/'evaluation_manifest.json').write_text(json.dumps({'task':TASK,'rules':RULES,'objectives':OBJECTIVES,'source_sha256':frozen,'profile':'general','max_requests':8,'max_tools':24,'max_seconds':600,'benchmark_code_sha256':fingerprint(__file__),'tool_controller_sha256':fingerprint(Path(__file__).resolve().parents[1]/'local_analyst/v3.py'),'cloud_requests':0},indent=2))
    import requests
    from local_analyst.providers import OllamaProvider
    from local_analyst.discovery import local_models
    installed=local_models();resident=requests.get('http://127.0.0.1:11434/api/ps',timeout=5).json()
    initial={m['name'] for m in resident.get('models',[])};runtime={'last_model':None,'unload_events':[],'initial_resident':sorted(initial)}
    def unload(model):
        if model and model not in initial:
            response=requests.post('http://127.0.0.1:11434/api/generate',json={'model':model,'keep_alive':0},timeout=60);response.raise_for_status()
            runtime['unload_events'].append({'model':model,'kind':'non-generating cache release'})
            (root/'resource_control.json').write_text(json.dumps(runtime,indent=2))
    def factory(model):
        if model not in installed:raise ValueError('Requested model unavailable; no download/fallback')
        if runtime['last_model']!=model:unload(runtime['last_model'])
        runtime['last_model']=model
        return OllamaProvider('http://127.0.0.1:11434/api/chat')
    common['provider_factory']=factory
    configs=[('A_single','single',{'analyst':'qwen3.8:27b'}),('B_comparison','comparison',{'a':'qwen3.8:27b','b':'qwen3:8b'}),('C_multi','multi',{'planner':'qwen3.8:27b','analyst':'qwen3.8:27b','reviewer':'qwen3.8:27b'})]
    for label,mode,models in configs:
        if any(model not in installed for model in models.values()):
            results[label]={'status':'NOT RUN','reason':'Selected second/local model unavailable; no replacement','api_requests':0,'api_usd':'0'}
            continue
        assert all(fingerprint(p)==h for p,h in frozen.items())
        result=execute(TASK,output=root/label,mode=mode,models=models,**common);results[label]=result
        (root/'summary.json').write_text(json.dumps(results,indent=2,default=str))
    unload(runtime['last_model'])
    (root/'resource_control.json').write_text(json.dumps(runtime,indent=2))
    return results

if __name__=='__main__':
    import argparse
    p=argparse.ArgumentParser();p.add_argument('--output',required=True);a=p.parse_args();results=run_suite(a.output)
    print(json.dumps({k:{'status':v['status'],'seconds':v.get('seconds'),'budget':v.get('budget')} for k,v in results.items()},indent=2))
