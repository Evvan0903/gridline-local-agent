"""Scripted generic SQL -> reviewed plan -> live XLSX -> PDF regression.

NOT a model evaluation. No reference answer is sent to any provider/network.
"""
from pathlib import Path
from decimal import Decimal as D,ROUND_HALF_UP
import json,socket,copy
from unittest.mock import patch
import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader
from local_analyst.v3 import execute
from local_analyst.financial_eligibility import validate_financial_rows
from local_analyst.verification import recalculate_workbook

OBJECTIVES=[{'id':k,'requirement':r} for k,r in [('plan','Plan the explicit source/eligibility work'),('eligibility','Payable-only coverage and currency-safe totals'),('calculation','Calculate governed base/scenario fees'),('excel','Generate a live parameter-driven workbook'),('report','Generate matching PDF with held/disputed audit')]]
CONTRACT={'id':'reviewed_rows','objective_id':'eligibility','kind':'financial_rowset','rule_ref':'Explicit fixture: ready payable, suspended held, challenged disputed; all other states unresolved','source_name':'records.csv','source_key':'id','status_column':'disposition','status_map':{'ready':'payable','suspended':'held','challenged':'disputed'},'require_payable_only':True,'totals':['base_usd','scenario_usd','delta_usd'],'monetary_columns':{'base_usd':{'fixed_currency':'USD'},'scenario_usd':{'fixed_currency':'USD'},'delta_usd':{'fixed_currency':'USD'},'base_local':{'currency_column':'currency'},'case_local':{'currency_column':'currency'}}}
RULES='Explicit synthetic regression: ready records only payable; suspended held and challenged disputed in separate audit. Base fee 10% net_amount USD; scenario parameter 12.5%. Local units multiply by source fx_rate. Round cents. Never sum unlike local currencies. Conditional test output, no payroll approval.'

def specification():
 return {'version':1,'rule_ref':RULES,'source_key':'id','parameters':{'base_rate':{'value':'.10','description':'Base fee rate'},'scenario_rate':{'value':'.125','description':'Editable scenario fee rate'}},'calculations':[
 {'id':'base_usd','expression':{'round':[{'mul':[{'field':'net_amount'},{'param':'base_rate'}]},2]}},
 {'id':'base_local','expression':{'round':[{'mul':[{'field':'base_usd'},{'field':'fx_rate'}]},2]}},
 {'id':'scenario_usd','expression':{'round':[{'mul':[{'field':'net_amount'},{'param':'scenario_rate'}]},2]}},
 {'id':'case_local','expression':{'round':[{'mul':[{'field':'scenario_usd'},{'field':'fx_rate'}]},2]}},
 {'id':'delta_usd','expression':{'sub':[{'field':'scenario_usd'},{'field':'base_usd'}]}}],'totals':['base_usd','scenario_usd','delta_usd']}

class ScriptedProvider:
 name='SCRIPTED_OFFLINE_NOT_A_MODEL'
 def __init__(self):self.requests=0
 def complete(self,model,history,tools):
  self.requests+=1;n=self.requests;results=[json.loads(m['content']).get('result',{}) for m in history if m.get('role')=='tool']
  table=next((x['tables'][0]['table_id'] for x in results if x.get('tables')),None)
  raw=next((x['query_id'] for x in results if x.get('columns')==['id','net_amount','fx_rate','currency']),None)
  plan=next((x['plan_id'] for x in results if x.get('plan_id')),None)
  artifact=next((x for x in results if x.get('kind')=='dynamic_financial_model'),None)
  calculated=next((x for x in results if 'base_usd' in x.get('columns',[])),None)
  def calls(*pairs):return {'role':'assistant','content':'Scripted regression step; no foundation model used.','tool_calls':[{'id':f'{n}_{i}','function':{'name':name,'arguments':args}} for i,(name,args) in enumerate(pairs)]}
  if n==1:msg=calls(('inspect_schema',{}),('set_task_plan',{'items':[{'id':x['id'],'steps':['Read source evidence','Apply explicit contract','Validate and deliver']} for x in OBJECTIVES]}))
  elif n==2:msg=calls(('query_sql',{'sql':f'SELECT id,net_amount,fx_rate,currency FROM "{table}" WHERE disposition=\'ready\' ORDER BY id'}))
  elif n==3:msg=calls(('define_financial_plan',{'spec':specification(),'query_id':raw}))
  elif n==4:msg=calls(('build_financial_model',{'plan_id':plan,'query_id':raw}))
  elif n==5:msg=calls(('query_sql',{'sql':f'SELECT id,currency,ROUND(net_amount*.10,2) AS base_usd,ROUND(net_amount*.10*fx_rate,2) AS base_local,ROUND(net_amount*.125,2) AS scenario_usd,ROUND(net_amount*.125*fx_rate,2) AS case_local,ROUND(net_amount*.125,2)-ROUND(net_amount*.10,2) AS delta_usd FROM "{table}" WHERE disposition=\'ready\' ORDER BY id'}),('validate_workbook',{'sha256':artifact['sha256'],'parameter_changes':{'B3':.20}}))
  elif n==6:
   rows=[dict(zip(calculated['columns'],r)) for r in calculated['rows']];totals={k:str(sum(D(str(row[k])) for row in rows)) for k in ('base_usd','scenario_usd','delta_usd')}
   msg=calls(('validate_financial_rows',{'query_id':calculated['query_id'],'contract_id':CONTRACT['id']}),('export_report',{'query_id':calculated['query_id'],'title':'Scripted financial regression','paragraphs':['Scripted offline engineering test. No Qwen or other model inference. USD totals: '+json.dumps(totals)+'. Local currency values remain separate. Suspended records are held; challenged records disputed, preserved in the eligibility audit. Source query '+calculated['query_id']]}))
  else:msg={'role':'assistant','content':'Scripted financial regression complete; payable-only results, live Excel and PDF generated. Held/disputed source records audited separately. Evidence '+calculated['query_id']}
  return msg,{'usage':{}}

def run_pipeline(root):
 root=Path(root);root.mkdir(parents=True,exist_ok=False);source=root/'records.csv'
 source.write_text('id,net_amount,fx_rate,currency,disposition\nalpha,70,1,USD,ready\nbeta,130,.9,EUR,ready\ngamma,40,2,JPY,suspended\ndelta,60,1.2,GBP,challenged\n')
 independent={}
 def verify(w,answer,events):
  frame=pd.read_csv(source);payable=frame[frame.disposition.eq('ready')];quant=lambda v:v.quantize(D('.01'),rounding=ROUND_HALF_UP)
  def source_oracle(rate):
   rows={}
   for row in payable.itertuples():
    net=D(str(row.net_amount));fx=D(str(row.fx_rate));base=quant(net*D('.10'));scenario=quant(net*D(rate))
    rows[row.id]={'base_usd':base,'base_local':quant(base*fx),'scenario_usd':scenario,'case_local':quant(scenario*fx),'delta_usd':scenario-base}
   return rows
  expected=source_oracle('.125')
  calc=[(q,w.query_result_frame(q)) for q in w.queries if {'base_usd','base_local','scenario_usd','case_local','delta_usd'}<=set(w.query_result_frame(q).columns)]
  math_ok=any(set(f.id)==set(expected) and len(f)==len(expected) and all(abs(D(str(row[k]))-expected[row['id']][k])<D('.005') for _,row in f.iterrows() for k in expected[row['id']]) for _,f in calc)
  row_checks=[validate_financial_rows(w,q,CONTRACT) for q,_ in calc];eligibility_ok=any(r['status']=='PASS' for r in row_checks)
  books=list((w.root/'outputs').rglob('*.xlsx'));excel_ok=False
  if len(books)==1:
   saved=load_workbook(books[0],data_only=True);formulas=load_workbook(books[0],data_only=False);rows=list(saved['Build'].values);data=[dict(zip(rows[0],r)) for r in rows[1:]]
   initial_ok=len(data)==len(expected) and set(x['id'] for x in data)==set(expected) and all(abs(D(str(row[k]))-expected[row['id']][k])<D('.005') for row in data for k in expected[row['id']])
   params=saved['Parameters'];cell=next(f'B{i}' for i in range(2,params.max_row+1) if params[f'A{i}'].value=='scenario_rate');changed=root/'changed_scenario.xlsx';recalc=recalculate_workbook(books[0],changed,{cell:.20});after=load_workbook(changed,data_only=True);values=list(after['Build'].values);changed_rows=[dict(zip(values[0],r)) for r in values[1:]]
   changed_expected=source_oracle('.20');changed_ok=len(changed_rows)==len(changed_expected) and set(row['id'] for row in changed_rows)==set(changed_expected) and all(abs(D(str(row[k]))-changed_expected[row['id']][k])<D('.005') for row in changed_rows for k in changed_expected[row['id']])
   totals={k:sum(D(str(row[k])) for row in data) for k in ('base_usd','scenario_usd','delta_usd')};summary={row[0]:D(str(row[1])) for row in list(saved['Summary'].values)[1:]};reconciled=all(abs(v-summary[k])<D('.005') for k,v in totals.items())
   excel_ok=initial_ok and changed_ok and reconciled and recalc['status']=='PASS' and any(c.data_type=='f' and 'Parameters' in str(c.value) for row in formulas['Build'] for c in row)
   independent.update(initial_values=data,changed_values=changed_rows,summary={k:str(v) for k,v in totals.items()},dynamic_scenario_pass=changed_ok,reconciliation_pass=reconciled,original_workbook=str(books[0]),changed_copy=str(changed))
  pdfs=list((w.root/'outputs').rglob('*.pdf'));pdf_ok=False
  if len(pdfs)==1:
   text=' '.join(p.extract_text() or '' for p in PdfReader(pdfs[0]).pages)
   pdf_ok='Scripted' in text and 'held' in text and all(str(sum(v[k] for v in expected.values())) in text for k in ('base_usd','scenario_usd','delta_usd')) and all(str(value) in text for row in expected.values() for value in row.values())
  checks={'plan':any(e.get('tool')=='set_task_plan' for e in events),'eligibility':eligibility_ok,'calculation':math_ok,'excel':excel_ok,'report':pdf_ok};independent.update(checks=checks,eligibility=row_checks)
  return [{'objective_id':k,'status':'PASS' if v else 'FAIL','reason':'Offline source/Decimal/rowset/live-XLSX/PDF verification; no foundation model score'} for k,v in checks.items()]
 provider=ScriptedProvider()
 with patch.object(socket.socket,'connect',side_effect=AssertionError('NO_NETWORK_SCRIPTED_PIPELINE')):
  result=execute('Scripted regression: generate dynamic Excel and PDF using the explicit governed financial rules and all original objectives.',[source],root/'run',models={'analyst':'SCRIPTED_OFFLINE_NOT_A_MODEL'},objectives=OBJECTIVES,rule_text=RULES,formats=['excel','pdf'],financial_contracts=[copy.deepcopy(CONTRACT)],provider_factory=lambda _:provider,verifier=verify)
 result_summary={'status':result['status'],'scripted_provider_turns':provider.requests,'real_model_requests':0,'api_requests':0,'cloud_cost_usd':'0','artifacts':result['artifacts'],'independent':independent}
 (root/'verification.json').write_text(json.dumps(result_summary,indent=2,default=str));return result_summary
if __name__=='__main__':
 import argparse
 p=argparse.ArgumentParser();p.add_argument('--output',required=True);args=p.parse_args();r=run_pipeline(args.output);print(json.dumps({'status':r['status'],'real_model_requests':0,'checks':r['independent'].get('checks')},indent=2));raise SystemExit(r['status']!='COMPLETED')
