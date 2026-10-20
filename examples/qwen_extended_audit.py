"""Offline audit after the once-only model attempt; never invokes inference."""
import json,hashlib
from pathlib import Path
from local_analyst.workspace import Workspace,fingerprint
from examples.qwen_extended_eval import OUT

def main():
 if not (OUT/'summary.json').exists():raise ValueError('Wait until real inference is terminal')
 manifest=json.loads((OUT/'evaluation_manifest.json').read_text());result=json.loads((OUT/'run/result.json').read_text());events=[json.loads(x) for x in (OUT/'run/events.jsonl').read_text().splitlines()]
 sources=list(manifest['sources']);w=Workspace(OUT/'independent_replay',restore=False);queries=[];quality=[]
 try:
  for i,p in enumerate(sources):w.load(p,'source'+str(i))
  for index,e in enumerate(events):
   if e.get('tool')=='query_sql' and e.get('result',{}).get('query_id'):
    old=e['result'];new=w.query(**e['arguments']);queries.append({'event_index':index,'query_id':old['query_id'],'sql':e['arguments']['sql'],'digest_match':new['result_digest']==old['result_digest'],'columns_match':new['columns']==old['columns'],'result_digest':old['result_digest']})
    evidence=w.queries.pop(new['query_id']);w.queries[old['query_id']]=evidence
   if e.get('tool')=='check_quality' and e.get('result',{}).get('status')!='error':
    new=w.quality_check(**e['arguments'],persist=False);quality.append({'arguments':e['arguments'],'affected_rows_match':new['affected_rows']==e['result']['affected_rows']})
 finally:w.db.close()
 from openpyxl import load_workbook
 from pypdf import PdfReader
 from examples.v3_benchmark import expected
 from decimal import Decimal as D
 oracle=expected(sources);books=[];pdfs=[]
 for artifact in result['artifacts']:
  p=Path(artifact['path']);assert fingerprint(p)==artifact['sha256']
  if p.suffix=='.xlsx':
   f=load_workbook(p,data_only=False);b=load_workbook(p,data_only=True)
   item={'path':str(p),'sheets':b.sheetnames,'formula_count':sum(c.data_type=='f' for s in f for row in s for c in row),'formula_errors':[{'sheet':s.title,'cell':c.coordinate,'error':c.value} for s in b for row in s for c in row if c.data_type=='e'],'parameters':list(b['Parameters'].values) if 'Parameters' in b else [],'summary':list(b['Summary'].values) if 'Summary' in b else []}
   if 'Build' in b:
    table=list(b['Build'].values);headers=list(table[0]);rows=[dict(zip(headers,row)) for row in table[1:] if row[0] is not None]
    item['build_rows']=rows
    try:
     item['detail_matches_source_oracle']=len(rows)==len(oracle) and {r['order_id'] for r in rows}==set(oracle) and all(abs(D(str(r[c]))-oracle[r['order_id']][c])<D('.005') for r in rows for c in oracle[r['order_id']])
     item['detail_usd_totals']={k:str(sum(D(str(r[k])) for r in rows)) for k in ('commission_usd','scenario_usd','delta_usd')}
    except Exception as e:item['detail_matches_source_oracle']=False;item['detail_error']=str(e)
   books.append(item)
  if p.suffix=='.pdf':
   reader=PdfReader(p);pdfs.append({'path':str(p),'pages':len(reader.pages),'text':'\n'.join(x.extract_text() or '' for x in reader.pages)})
 protected=json.loads((OUT.parent/'extended_preflight/protected_files.json').read_text());changed=[p for p,h in protected.items() if not Path(p).exists() or fingerprint(p)!=h]
 actions=[{'event_index':i,**{k:e.get(k) for k in ('kind','tool','arguments','result','error','budget')}} for i,e in enumerate(events) if e.get('kind') in {'tool','tool_error'}]
 for name in ('query_sql','validate_sql','define_financial_plan','set_task_plan'):
  (OUT/(name+'_attempts.json')).write_text(json.dumps([a for a in actions if a['tool']==name],indent=2,default=str))
 report={'inference_requests_in_audit':0,'paid_requests':0,'sql_replays':queries,'quality_replays':quality,'all_sql_digests_match':all(q['digest_match'] and q['columns_match'] for q in queries),'workbooks':books,'pdfs':pdfs,'protected_files':len(protected),'historical_or_source_changes':changed,'model_result_sha256':fingerprint(OUT/'run/result.json'),'frozen_code_unchanged':all(fingerprint(Path(p))==h for p,h in manifest['code_sha256'].items()),'checks':result['validation']['checks']}
 (OUT/'independent_audit.json').write_text(json.dumps(report,indent=2,default=str));print(json.dumps({k:report[k] for k in ('all_sql_digests_match','protected_files','historical_or_source_changes','frozen_code_unchanged')}))
if __name__=='__main__':main()
