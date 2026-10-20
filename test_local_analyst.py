from decimal import Decimal
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch, Mock

import pandas as pd

from local_analyst import Workspace
from local_analyst.finance import marginal_tiers, ramp_credit, window_index, hierarchy_paths, convert_currency
from local_analyst.modeling import run_plan, scenario, validate_plan, inspect_rules
from local_analyst.exporting import export_query
from local_analyst.agent import AnalystAgent


class CoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.workspace = Workspace(self.root / "workspace")
        self.addCleanup(self.workspace.db.close)
        self.workspace.register(pd.DataFrame({"key": [1, 1, 2, None], "amount": [10, 20, 30, 40],
                                               "__source_row": [2, 3, 4, 5]}), "orders", {"kind": "fixture"})
        self.workspace.register(pd.DataFrame({"ref": [1, 1, 3]}), "customers", {"kind": "fixture"})

    def test_stable_csv_loading_and_refresh(self):
        path = self.root / "sales.csv"
        path.write_text("id,amount,country\na,1,NA\nb,2,US\n")
        first = self.workspace.load(path, "sales")["tables"][0]
        self.assertEqual(self.workspace.frames[first["table_id"]].country.iloc[0], "NA")
        path.write_text("id,amount,country\na,10,NA\nb,20,US\nc,30,MX\n")
        second = self.workspace.load(path, "sales")["tables"][0]
        self.assertEqual(first["table_id"], second["table_id"])
        self.assertNotEqual(first["source_sha256"], second["source_sha256"])
        restored = Workspace(self.root / "workspace")
        self.addCleanup(restored.db.close)
        self.assertEqual(restored.schema(second["table_id"])["rows"], 3)

    def test_query_parameters_limit_and_audit(self):
        result = self.workspace.query("SELECT key, sum(amount) AS total FROM orders WHERE amount > $min GROUP BY key ORDER BY total", {"min": 15}, 2)
        self.assertEqual(result["rows"], [[1.0, 20], [2.0, 30]])
        self.assertTrue(result["truncated"])
        evidence = json.loads(self.workspace.audit_path.read_text().splitlines()[-1])
        self.assertEqual(evidence["parameters"], {"min": 15})
        self.assertEqual(evidence["event_id"], result["query_id"])

    def test_readonly_security(self):
        attacks = ["DELETE FROM orders", "SELECT 1; SELECT 2", "ATTACH '/tmp/a.db'", "INSTALL httpfs",
                   "COPY orders TO '/tmp/unsafe.csv'", "SELECT * FROM read_csv('/etc/passwd')",
                   "SELECT * FROM glob('/Users/*')", "SELECT getenv('HOME')", "SELECT * FROM duckdb_settings()",
                   "SELECT * FROM orders CROSS JOIN customers", "SELECT * FROM orders NATURAL JOIN customers",
                   "WITH RECURSIVE x AS (SELECT 1 UNION ALL SELECT * FROM x) SELECT * FROM x",
                   "SELECT * FROM orders INTO new_table", "SELECT * FROM 'file.csv'"]
        for sql in attacks:
            with self.subTest(sql=sql), self.assertRaises(Exception):
                self.workspace.query(sql)
        self.assertEqual(self.workspace.query("SELECT count(*) FROM orders")["rows"], [[4]])

    def test_join_cardinality_nulls_and_quantified_ledger(self):
        audit = self.workspace.join_check("orders", "customers", ["key"], ["ref"])
        self.assertEqual(audit["cardinality"], "many_to_many")
        self.assertEqual(audit["inner_join_rows"], 4)
        self.assertEqual(audit["left"]["null_rows"], 1)
        queried = self.workspace.query("SELECT o.key,c.ref FROM orders o JOIN customers c ON o.key=c.ref")
        self.assertIn("multiplication", queried["warnings"][0]["warning"])
        ledger = self.workspace.quality_check("orders", "duplicates", ["key"], "amount")
        self.assertEqual(ledger["affected_rows"], 2)
        self.assertEqual(ledger["exposed_amount"], "30")
        self.assertEqual(pd.read_csv(ledger["ledger_csv"])["__source_row"].tolist(), [2, 3])

    def test_reconciliation_uses_scalar_evidence(self):
        result = self.workspace.reconcile("SELECT sum(amount) FROM orders", "SELECT 100")
        self.assertTrue(result["passed"])
        self.assertEqual(result["difference"], "0")
        with self.assertRaises(ValueError):
            self.workspace.reconcile("SELECT NULL", "SELECT 0")

    def test_export_roundtrip_and_lineage(self):
        result = self.workspace.query("SELECT key, amount, __source_row FROM orders ORDER BY amount")
        files = export_query(self.workspace, result["query_id"], "checked", pdf=True)
        self.assertEqual(len(pd.read_csv(files["csv"])), 4)
        import openpyxl
        book = openpyxl.load_workbook(files["excel"]["path"], data_only=True)
        self.assertEqual(book["Result"].max_row, 5)
        self.assertIn("Provenance", book.sheetnames)
        book.close()
        from pypdf import PdfReader
        self.assertIn(result["query_id"], PdfReader(files["pdf"]["path"]).pages[0].extract_text())

    def test_agent_loop_executes_tools_and_handles_errors(self):
        agent = AnalystAgent(self.workspace)
        responses = [
            {"role":"assistant","content":"","tool_calls":[{"function":{"name":"query_sql","arguments":{"sql":"DELETE FROM orders"}}}]},
            {"role":"assistant","content":"","tool_calls":[{"function":{"name":"query_sql","arguments":json.dumps({"sql":"SELECT count(*) AS n FROM orders"})}}]},
            {"role":"assistant","content":"There are 4 rows."},
        ]
        def response(*args, **kwargs):
            value = Mock(); value.json.return_value = {"message": responses.pop(0)}; return value
        with patch("local_analyst.agent.requests.post", side_effect=response):
            answer = agent.ask("Count rows")
        self.assertIn("4", answer["answer"])
        self.assertIn("error", answer["trace"][0]["result"])
        self.assertEqual(answer["trace"][1]["result"]["rows"], [[4]])


def simple_plan(table):
    return {"bindings": {"invoices": table}, "parameters": {
        "rate": {"type":"number", "value":"0.1"}, "threshold": {"type":"number", "value":"100"},
        "multiplier": {"type":"number", "value":"2"}, "asof": {"type":"date", "value":"2026-06-30"}},
        "assumptions": [{"status":"explicit_user_choice","description":"Tiered collection fee on paid invoices by asof date"}],
        "steps": [
            {"id":"eligible","kind":"sql","rule_ref":"User supplied collections policy 1","sql":"SELECT invoice, rep, CAST(gross AS DECIMAL(18,2)) AS gross, CAST(fx AS DECIMAL(18,4)) AS fx FROM {{invoices}} WHERE paid = true AND CAST(due AS DATE) <= CAST($asof AS DATE)"},
            {"id":"fees","kind":"tiers","rule_ref":"Collections policy 2","source":"eligible","amount_column":"gross","rate_parameter":"rate","thresholds":["threshold"],"multipliers":[1,"multiplier"],"output_column":"fee"},
            {"id":"converted","kind":"currency","rule_ref":"Collections policy 3","source":"fees","amount_column":"fee","rate_column":"fx","output_column":"local_fee"},
        ]}


class FinancialReuseTests(unittest.TestCase):
    def test_primitives_against_hand_calculations(self):
        self.assertEqual(marginal_tiers(45000,.16,[30000,40000],[1,1.2,1.3]), Decimal("7760.00"))
        self.assertEqual(marginal_tiers(-100,.1), Decimal("-10.0"))
        self.assertEqual(window_index("2026-06-10","2026-05-12",30),0)
        self.assertEqual(window_index("2026-06-11","2026-05-12",30),1)
        self.assertEqual(window_index("2026-05-01","2026-05-12",30,True),0)
        ramp = ramp_credit("2026-06-01",30000,"2026-05-18","2026-05-18",None,[1,.5])
        self.assertEqual(ramp["floor"],Decimal(23000))  # 16 full days + 14 half days
        self.assertEqual(convert_currency("12.34","1.25"),Decimal("15.4250"))

    def test_hierarchy_cycle_and_unknown_parent(self):
        graph = pd.DataFrame({"person":["a","b","c"],"boss":["b","c",None]})
        paths,warnings = hierarchy_paths(graph,"person","boss")
        self.assertEqual(paths["a"],["b","c"])
        self.assertFalse(warnings)
        graph.loc[2,"boss"]="a"
        with self.assertRaises(ValueError): hierarchy_paths(graph,"person","boss")
        graph.loc[2,"boss"]="external"
        self.assertTrue(hierarchy_paths(graph,"person","boss")[1])

    def test_different_schema_full_scenario_dependency_and_export(self):
        with tempfile.TemporaryDirectory() as root:
            w=Workspace(root);self.addCleanup(w.db.close)
            f=pd.DataFrame({"invoice":["I1","I2","I3"],"rep":["P","Q","P"],"gross":[150,80,999],"fx":[1.25,1.25,1.25],"paid":[True,True,False],"due":["2026-06-02"]*3})
            w.register(f,"invoices",{"kind":"fixture"})
            p=simple_plan("invoices")
            result=scenario(w,p,{"rate":"0.2"},"converted","local_fee")
            self.assertEqual(Decimal(result["baseline"]),Decimal(35))
            self.assertEqual(Decimal(result["scenario"]),Decimal(70))
            self.assertEqual(Decimal(result["delta"]),Decimal(35))
            self.assertEqual(w.frames["model_baseline_converted"].fee.tolist(),[Decimal(20),Decimal(8)])

    def test_missing_ambiguous_and_unsupported_rules_block(self):
        p=simple_plan("anything")
        p["parameters"]["rate"].pop("value")
        self.assertEqual(validate_plan(p)["status"],"blocked")
        p=simple_plan("anything");p["assumptions"].append({"status":"unresolved","description":"unclear refund treatment"})
        self.assertEqual(validate_plan(p)["status"],"blocked")
        p=simple_plan("anything");p["steps"][0]["kind"]="execute_python"
        self.assertEqual(validate_plan(p)["status"],"blocked")
        with tempfile.TemporaryDirectory() as root:
            path=Path(root)/"rules.md";path.write_text("Pay a reasonable bonus. Maybe use sales?")
            self.assertTrue(inspect_rules(path)["warnings"])


if __name__ == "__main__":
    unittest.main()
